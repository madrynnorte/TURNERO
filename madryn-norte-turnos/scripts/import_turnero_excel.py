from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX


TZ = ZoneInfo("America/Argentina/Buenos_Aires")
COURTS = {
    "main": ("court-main", range(2, 9)),
    "multi": ("court-multi", range(11, 18)),
}
DEFAULT_COLORS = {
    "type-rental": "#DCEBF9",
    "type-fixed": "#D9F0E2",
    "type-birthday": "#EEE6FB",
    "type-tournament": "#D9F0E2",
    "type-training": "#DFF3F1",
}


def clean(value: object) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip())
    text = re.sub(r"\s+fijo\s*$", "", text, flags=re.I)
    aliases = {
        "hockey mn": "Hockey MN",
        "entrenamiento marcos mn": "Entrenamiento Marcos MN",
        "equipo femenino marcos mn": "Equipo Femenino Marcos MN",
        "futbol nuria mn": "Fútbol Nuria MN",
        "fut adol mn": "Fútbol adolescentes MN",
        "futbol adol mn": "Fútbol adolescentes MN",
        "futbol adol. mn": "Fútbol adolescentes MN",
        "futbol adol mn ": "Fútbol adolescentes MN",
        "basquet fcn": "Básquet FCN",
        "arqueritos tomi": "Arqueritos Tomi",
        "entrenamiento arqueros tomi": "Entrenamiento arqueros Tomi",
        "lisisadas fc": "Lisiadas FC",
        "lisiadas fc": "Lisiadas FC",
        "torneo masc": "Torneo masculino",
        "torneo +30": "Torneo +30",
        "torneo polis": "Torneo Polis",
        "silvina hockey": "Silvina hockey",
        "cervecero": "Cervecero",
        "jose luis": "José Luis",
        "naim": "Naim",
        "andy": "Andy",
    }
    return aliases.get(text.casefold(), text[:1].upper() + text[1:])


def plain(value: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", value.casefold()) if unicodedata.category(ch) != "Mn")


def classify(name: str, raw: str) -> tuple[str, str]:
    key = plain(f"{name} {raw}")
    if "cumple" in key:
        return "type-birthday", "por_encuentro"
    if "torneo" in key:
        return "type-tournament", "por_encuentro"
    if any(word in key for word in ("entren", "hockey", "basquet", "futbol", "arquero", "equipo femenino", "adol mn")):
        return "type-training", "mensual"
    if "fijo" in key or plain(name) in {"banfield", "lisiadas fc", "andy", "mauricio"}:
        return "type-fixed", "mensual"
    return "type-rental", "por_encuentro"


def color_for(cell, booking_type: str) -> str:
    color = cell.fill.fgColor
    if color.type == "rgb" and color.rgb:
        raw = color.rgb[-6:]
        if raw != "000000":
            return f"#{raw.upper()}"
    if color.type == "indexed" and color.indexed is not None and color.indexed < len(COLOR_INDEX):
        raw = COLOR_INDEX[color.indexed][-6:]
        if raw not in {"000000", "FFFFFF"}:
            return f"#{raw.upper()}"
    return DEFAULT_COLORS[booking_type]


def sql(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def iso(local_dt: datetime) -> str:
    return local_dt.replace(tzinfo=TZ).astimezone(ZoneInfo("UTC")).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def merged_lookup(sheet) -> dict[tuple[int, int], tuple[object, object]]:
    result: dict[tuple[int, int], tuple[object, object]] = {}
    for merged in sheet.merged_cells.ranges:
        source = sheet.cell(merged.min_row, merged.min_col)
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                result[(row, column)] = (source.value, source)
    return result


def parse_bookings(path: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    bookings: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        merged = merged_lookup(sheet)
        for _, (court_id, columns) in COURTS.items():
            for column in columns:
                date_value = sheet.cell(3, column).value
                if not isinstance(date_value, datetime):
                    continue
                slots: list[tuple[int, str, str, object]] = []
                for row in range(4, 34):
                    value, source_cell = merged.get((row, column), (sheet.cell(row, column).value, sheet.cell(row, column)))
                    if value is None or not str(value).strip():
                        continue
                    raw = re.sub(r"\s+", " ", str(value).strip())
                    slots.append((row, clean(raw), raw, source_cell))
                index = 0
                while index < len(slots):
                    start_row, name, raw, source_cell = slots[index]
                    end_row = start_row + 1
                    index += 1
                    while index < len(slots) and slots[index][0] == end_row and slots[index][1].casefold() == name.casefold():
                        end_row += 1
                        index += 1
                    start_minutes = 9 * 60 + (start_row - 4) * 30
                    end_minutes = 9 * 60 + (end_row - 4) * 30
                    start_local = date_value.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(minutes=start_minutes)
                    end_local = date_value.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(minutes=end_minutes)
                    booking_type, payment_mode = classify(name, raw)
                    start_iso, end_iso = iso(start_local), iso(end_local)
                    digest = hashlib.sha1(f"{court_id}|{start_iso}|{end_iso}|{name}".encode("utf-8")).hexdigest()[:20]
                    series = None
                    if payment_mode == "mensual":
                        series = "series-import-" + hashlib.sha1(f"{court_id}|{plain(name)}".encode("utf-8")).hexdigest()[:14]
                    price = 40000 if booking_type == "type-rental" and court_id == "court-main" else 25000 if booking_type == "type-rental" else 0
                    bookings.append({
                        "id": f"turnero-{digest}", "name": name, "court_id": court_id, "booking_type": booking_type,
                        "starts_at": start_iso, "ends_at": end_iso, "color": color_for(source_cell, booking_type),
                        "payment_mode": payment_mode, "list_price": price, "final_price": price, "series_id": series,
                        "notes": f"Importado de {path.name} · {sheet.title} {source_cell.coordinate}",
                    })
    unique = {booking["id"]: booking for booking in bookings}
    return sorted(unique.values(), key=lambda item: (item["starts_at"], item["court_id"], item["name"]))


def write_sql(bookings: list[dict[str, object]], output: Path) -> None:
    lines = [
        "PRAGMA foreign_keys = ON;",
        "UPDATE courts SET name = 'Cancha 7 sintético' WHERE id = 'court-main';",
        "UPDATE courts SET name = 'Cancha multideporte (fútbol 5)' WHERE id = 'court-multi';",
        "DELETE FROM reservation_slots WHERE reservation_id IN (SELECT id FROM reservations WHERE source = 'excel');",
        "DELETE FROM reservations WHERE source = 'excel';",
    ]
    reservation_rows = []
    slot_rows = []
    for booking in bookings:
        reservation_rows.append("(" + ",".join(sql(value) for value in (
            booking["id"], booking["name"], "", booking["court_id"], booking["booking_type"], booking["starts_at"], booking["ends_at"],
            "confirmado", booking["color"], booking["notes"], booking["payment_mode"], booking["list_price"], "none", 0, "", booking["final_price"],
            None, booking["series_id"], "excel", 1, "user-admin",
        )) + ")")
        slot_time = datetime.fromisoformat(str(booking["starts_at"]).replace("Z", "+00:00"))
        end_time = datetime.fromisoformat(str(booking["ends_at"]).replace("Z", "+00:00"))
        while slot_time < end_time:
            slot_iso = slot_time.isoformat(timespec="milliseconds").replace("+00:00", "Z")
            slot_id = "slot-" + hashlib.sha1(f"{booking['id']}|{slot_iso}".encode("utf-8")).hexdigest()[:24]
            slot_rows.append("(" + ",".join(sql(value) for value in (slot_id, booking["id"], booking["court_id"], slot_iso)) + ")")
            slot_time += timedelta(minutes=30)
    for offset in range(0, len(reservation_rows), 75):
        lines.append("INSERT INTO reservations (id,client_name,phone,court_id,booking_type_id,starts_at,ends_at,status,color,notes,payment_mode,list_price,discount_type,discount_value,discount_reason,final_price,responsible_id,series_id,source,version,created_by) VALUES\n" + ",\n".join(reservation_rows[offset:offset + 75]) + ";")
    for offset in range(0, len(slot_rows), 150):
        lines.append("INSERT INTO reservation_slots (id,reservation_id,court_id,slot_at) VALUES\n" + ",\n".join(slot_rows[offset:offset + 150]) + ";")
    lines.append("")
    output.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    bookings = parse_bookings(args.source)
    write_sql(bookings, args.output)
    by_court = {court: sum(1 for item in bookings if item["court_id"] == court) for court in ("court-main", "court-multi")}
    print(json.dumps({
        "reservations": len(bookings), "byCourt": by_court,
        "from": bookings[0]["starts_at"] if bookings else None,
        "to": bookings[-1]["ends_at"] if bookings else None,
        "sql": str(args.output),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
